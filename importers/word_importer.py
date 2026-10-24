"""واردسازی فایل‌های Word ساختار سازمانی و رفاهی"""
import logging
from pathlib import Path
from docx import Document
from sqlalchemy.orm import Session
from services.category_service import get_or_create_category
from services.question_service import create_question, get_or_create_service
from database.models import ImportJob, ContactPoint
from datetime import datetime, timezone
import re

logger = logging.getLogger(__name__)


def _clean(text) -> str:
    if text is None:
        return ""
    return str(text).replace("\xa0", " ").strip()


def import_word(session: Session, file_path: Path, root_name: str = None) -> dict:
    """
    واردسازی جدول استاندارد ۷ ستونه فایل‌های Word.
    """
    job = ImportJob(source_file=str(file_path.name), status="running")
    session.add(job)
    session.flush()

    report = {
        "file": file_path.name,
        "total": 0,
        "imported": 0,
        "skipped": 0,
        "failed": 0,
        "errors": [],
        "categories": set(),
        "services": set(),
    }

    try:
        doc = Document(file_path)
        if not doc.tables:
            job.status = "failed"
            job.report = "جدولی در فایل یافت نشد"
            job.finished_at = datetime.now(timezone.utc)
            report["errors"].append(job.report)
            return report

        table = doc.tables[0]
        rows = table.rows
        if len(rows) < 2:
            job.status = "failed"
            job.report = "جدول داده ندارد"
            job.finished_at = datetime.now(timezone.utc)
            return report

        # تشخیص هدر
        header_cells = [_clean(c.text) for c in rows[0].cells]
        # انتظار: سطح اول، سطح دوم، سطح سوم، عنوان خدمت، پرسش، پاسخ، کلمات کلیدی

        # ریشه اصلی در صورت نیاز
        root = None
        if root_name:
            root = get_or_create_category(
                session, root_name, parent_id=None, source_file=file_path.name
            )

        for idx, row in enumerate(rows[1:], start=2):
            report["total"] += 1
            try:
                cells = [_clean(c.text) for c in row.cells]
                # اطمینان از حداقل ۷ ستون
                while len(cells) < 7:
                    cells.append("")

                level1 = cells[0]
                level2 = cells[1]
                level3 = cells[2]
                service_title = cells[3]
                question = cells[4]
                answer = cells[5]
                keywords_raw = cells[6]

                if not question or not answer:
                    report["skipped"] += 1
                    report["errors"].append(f"ردیف {idx}: پرسش یا پاسخ خالی")
                    continue

                # ساخت سلسله‌مراتب
                parent_id = root.id if root else None
                cat = None

                for level_name in (level1, level2, level3):
                    if level_name:
                        cat = get_or_create_category(
                            session,
                            level_name,
                            parent_id=parent_id,
                            source_file=file_path.name,
                        )
                        parent_id = cat.id
                        report["categories"].add(level_name)

                if not cat and root:
                    cat = root
                elif not cat:
                    # بدون سطح → دسته‌بندی پیش‌فرض
                    cat = get_or_create_category(
                        session, "سایر", parent_id=None, source_file=file_path.name
                    )

                # خدمت
                service = None
                if service_title:
                    service = get_or_create_service(
                        session, service_title, cat.id, source_file=file_path.name
                    )
                    report["services"].add(service_title)

                # کلمات کلیدی
                keywords = []
                if keywords_raw:
                    for part in re.split(r"[,،\n|/]+", keywords_raw):
                        part = part.strip()
                        if part:
                            keywords.append(part)

                create_question(
                    session,
                    question_text=question,
                    answer_text=answer,
                    category_id=cat.id,
                    service_id=service.id if service else None,
                    keywords=keywords,
                    source_file=file_path.name,
                    source_row=idx,
                )
                report["imported"] += 1

                # استخراج شماره‌های داخلی از پاسخ (ساده)
                _extract_contacts(session, answer, level1 or level2 or "")

            except Exception as e:
                report["failed"] += 1
                report["errors"].append(f"ردیف {idx}: {str(e)}")
                logger.exception("Word import row error")

        job.total_rows = report["total"]
        job.imported = report["imported"]
        job.skipped = report["skipped"]
        job.failed = report["failed"]
        job.status = "success" if report["failed"] == 0 else "partial"
        job.report = (
            f"وارد شده: {report['imported']} | رد شده: {report['skipped']} | "
            f"خطا: {report['failed']} | دسته‌ها: {len(report['categories'])} | "
            f"خدمات: {len(report['services'])}"
        )
        job.finished_at = datetime.now(timezone.utc)
        session.flush()

    except Exception as e:
        job.status = "failed"
        job.report = str(e)
        job.finished_at = datetime.now(timezone.utc)
        report["errors"].append(str(e))
        logger.exception("Word import failed")

    report["categories"] = list(report["categories"])
    report["services"] = list(report["services"])
    return report


def _extract_contacts(session: Session, text: str, unit: str):
    """استخراج ساده شماره‌های داخلی از متن پاسخ"""
    # الگوی داخلی مثل ۲۱۴۰ یا 2140
    matches = re.findall(r"(?:داخلی|داخلی\s*)[^\d]*(\d{3,5})", text)
    for m in matches:
        existing = (
            session.query(ContactPoint)
            .filter(ContactPoint.internal == m)
            .first()
        )
        if not existing:
            session.add(
                ContactPoint(
                    name=f"داخلی {m}",
                    unit=unit[:100],
                    internal=m,
                    description=text[:300],
                )
            )
            session.flush()
