"""واردسازی فایل اکسل بیمه تکمیلی"""
import logging
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from services.category_service import get_or_create_category
from services.question_service import create_question
from database.models import ImportJob
from datetime import datetime, timezone

logger = logging.getLogger(__name__)


def import_excel(session: Session, file_path: Path) -> dict:
    """
    واردسازی فایل اکسل بیمه تکمیلی.
    برمی‌گرداند دیکشنری گزارش.
    """
    job = ImportJob(source_file=str(file_path.name), status="running")
    session.add(job)
    session.flush()

    report = {
        "file": file_path.name,
        "total": 0,
        "imported": 0,
        "skipped": 0,
        "failed": 0,
        "errors": [],
        "categories": set(),
    }

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            job.status = "failed"
            job.report = "فایل خالی است"
            job.finished_at = datetime.now(timezone.utc)
            return report

        header = [str(c or "").strip() for c in rows[0]]
        # نگاشت ستون‌ها
        col_map = {}
        for i, h in enumerate(header):
            h_norm = h.replace("\u200c", "").strip()
            if "دسته" in h_norm:
                col_map["category"] = i
            elif "پرسش" in h_norm:
                col_map["question"] = i
            elif "پاسخ" in h_norm:
                col_map["answer"] = i
            elif "زیرموضوع" in h_norm or "موضوع" in h_norm:
                col_map["subtopic"] = i

        if "question" not in col_map or "answer" not in col_map:
            job.status = "failed"
            job.report = f"ستون‌های ضروری یافت نشد. هدر: {header}"
            job.finished_at = datetime.now(timezone.utc)
            report["errors"].append(job.report)
            return report

        # ریشه بیمه تکمیلی
        root = get_or_create_category(
            session, "بیمه تکمیلی", parent_id=None, source_file=file_path.name, sort_order=1
        )

        for idx, row in enumerate(rows[1:], start=2):
            report["total"] += 1
            try:
                cat_name = str(row[col_map.get("category", 0)] or "").strip()
                question = str(row[col_map["question"]] or "").strip()
                answer = str(row[col_map["answer"]] or "").strip()
                subtopic = ""
                if "subtopic" in col_map:
                    subtopic = str(row[col_map["subtopic"]] or "").strip()

                if not question or not answer:
                    report["skipped"] += 1
                    report["errors"].append(f"ردیف {idx}: پرسش یا پاسخ خالی")
                    continue

                # دسته‌بندی سطح دوم
                if cat_name:
                    cat = get_or_create_category(
                        session,
                        cat_name,
                        parent_id=root.id,
                        source_file=file_path.name,
                    )
                    report["categories"].add(cat_name)
                else:
                    cat = root

                keywords = [subtopic] if subtopic else []

                create_question(
                    session,
                    question_text=question,
                    answer_text=answer,
                    category_id=cat.id,
                    keywords=keywords,
                    source_file=file_path.name,
                    source_row=idx,
                )
                report["imported"] += 1

            except Exception as e:
                report["failed"] += 1
                report["errors"].append(f"ردیف {idx}: {str(e)}")
                logger.exception("Excel import row error")

        job.total_rows = report["total"]
        job.imported = report["imported"]
        job.skipped = report["skipped"]
        job.failed = report["failed"]
        job.status = "success" if report["failed"] == 0 else "partial"
        job.report = (
            f"وارد شده: {report['imported']} | رد شده: {report['skipped']} | "
            f"خطا: {report['failed']} | دسته‌ها: {len(report['categories'])}"
        )
        job.finished_at = datetime.now(timezone.utc)
        session.flush()

    except Exception as e:
        job.status = "failed"
        job.report = str(e)
        job.finished_at = datetime.now(timezone.utc)
        report["errors"].append(str(e))
        logger.exception("Excel import failed")

    report["categories"] = list(report["categories"])
    return report
